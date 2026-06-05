"""
文档生成器
支持 PDF 和 Word 格式的会议纪要导出
"""
import os
from datetime import datetime
from typing import Dict, Any, List

class DocumentGenerator:
    """会议纪要文档生成器"""

    def __init__(self, output_dir: str = "./outputs"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    async def generate_pdf(self, meeting_data: Dict[str, Any], filename: str) -> str:
        """
        生成 PDF 格式会议纪要

        Args:
            meeting_data: 会议数据
            filename: 输出文件名

        Returns:
            生成的 PDF 文件路径
        """
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib import colors

            file_path = os.path.join(self.output_dir, filename)

            doc = SimpleDocTemplate(file_path, pagesize=A4)
            styles = getSampleStyleSheet()

            # 创建自定义样式
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                spaceAfter=30,
                alignment=1  # 居中
            )

            # 构建文档内容
            story = []

            # 标题
            story.append(Paragraph(meeting_data.get("title", "会议纪要"), title_style))
            story.append(Spacer(1, 0.5*cm))

            # 会议信息
            info_text = f"""会议日期: {meeting_data.get('date', datetime.now().strftime('%Y-%m-%d %H:%M'))}<br/>
            参会人数: {len(meeting_data.get('speakers', []))} 人<br/>
            发言记录: {len(meeting_data.get('transcripts', []))} 条"""
            story.append(Paragraph(info_text, styles['Normal']))
            story.append(Spacer(1, 1*cm))

            # 参会人员
            if meeting_data.get("speakers"):
                story.append(Paragraph("参会人员", styles['Heading2']))
                speaker_list = "<br/>".join([
                    f"- {s.get('name', s.get('speaker_id', 'Unknown'))}"
                    for s in meeting_data["speakers"]
                ])
                story.append(Paragraph(speaker_list, styles['Normal']))
                story.append(Spacer(1, 0.5*cm))

            # 发言记录
            story.append(Paragraph("发言记录", styles['Heading2']))

            for i, transcript in enumerate(meeting_data.get("transcripts", []), 1):
                speaker = transcript.get("speaker", "Unknown")
                text = transcript.get("text", "")
                translation = transcript.get("translation", "")

                content = f"""<b>{i}. {speaker}</b><br/>
                原文: {text}<br/>
                翻译: {translation if translation else '（无）'}"""
                story.append(Paragraph(content, styles['Normal']))
                story.append(Spacer(1, 0.3*cm))

            # 生成 PDF
            doc.build(story)
            return file_path

        except ImportError:
            # 如果 reportlab 未安装，使用模拟
            file_path = os.path.join(self.output_dir, filename)
            with open(file_path.replace(".pdf", "_mock.txt"), "w", encoding="utf-8") as f:
                f.write(f"会议纪要 - {filename}\n")
                f.write(f"生成时间: {datetime.now().isoformat()}\n\n")
                f.write(f"标题: {meeting_data.get('title')}\n")
                f.write(f"发言记录数: {len(meeting_data.get('transcripts', []))}\n")
            return file_path.replace(".pdf", "_mock.txt")

    async def generate_docx(self, meeting_data: Dict[str, Any], filename: str) -> str:
        """
        生成 Word 格式会议纪要

        Args:
            meeting_data: 会议数据
            filename: 输出文件名

        Returns:
            生成的 Word 文件路径
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            file_path = os.path.join(self.output_dir, filename)
            doc = Document()

            # 标题
            title = doc.add_heading(meeting_data.get("title", "会议纪要"), 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 会议信息
            doc.add_paragraph()
            info = doc.add_paragraph()
            info.add_run("会议信息").bold = True
            info.add_run(f"\n日期: {meeting_data.get('date', datetime.now().strftime('%Y-%m-%d %H:%M'))}")
            info.add_run(f"\n参会人数: {len(meeting_data.get('speakers', []))} 人")

            # 参会人员
            if meeting_data.get("speakers"):
                doc.add_heading("参会人员", level=1)
                for speaker in meeting_data["speakers"]:
                    doc.add_paragraph(f"- {speaker.get('name', speaker.get('speaker_id', 'Unknown'))}")

            # 发言记录
            doc.add_heading("发言记录", level=1)

            for i, transcript in enumerate(meeting_data.get("transcripts", []), 1):
                speaker = transcript.get("speaker", "Unknown")
                text = transcript.get("text", "")
                translation = transcript.get("translation", "")

                p = doc.add_paragraph()
                p.add_run(f"{i}. {speaker}: ").bold = True
                p.add_run(f"{text}")
                if translation:
                    p.add_run(f"\n   翻译: {translation}")

            # 保存
            doc.save(file_path)
            return file_path

        except ImportError:
            # 如果 python-docx 未安装，使用模拟
            file_path = os.path.join(self.output_dir, filename)
            with open(file_path.replace(".docx", "_mock.txt"), "w", encoding="utf-8") as f:
                f.write(f"会议纪要 - {filename}\n")
                f.write(f"生成时间: {datetime.now().isoformat()}\n\n")
                f.write(f"标题: {meeting_data.get('title')}\n")
                f.write(f"发言记录数: {len(meeting_data.get('transcripts', []))}\n")
            return file_path.replace(".docx", "_mock.txt")